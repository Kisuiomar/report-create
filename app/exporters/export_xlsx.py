"""
XLSX Exporter for DossierReport.
Creates an Excel workbook with 4 structured and styled tabs:
«Анкета», «Контакты и адреса», «Родственные связи», «Места работы».
"""

from pathlib import Path
from typing import Union

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.schemas import DossierReport


def _style_header_row(ws, row_idx: int, cols_count: int, title: str = None, bg_hex: str = "1E293B"):
    """Стилизация строки заголовка таблицы."""
    fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
    font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    border = Border(
        left=Side(style="thin", color="94A3B8"),
        right=Side(style="thin", color="94A3B8"),
        top=Side(style="thin", color="94A3B8"),
        bottom=Side(style="thin", color="94A3B8")
    )
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col in range(1, cols_count + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = align
    ws.row_dimensions[row_idx].height = 26


def _auto_fit_columns(ws, max_width: int = 50):
    """Автоматическая подгонка ширины колонок."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            # Игнорируем длинные объединенные строки описаний при расчете
            if len(val) > 80:
                continue
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), max_width)


def export_xlsx(dossier: DossierReport, output_path: Union[str, Path]) -> Path:
    """
    Формирует Excel-файл со структурированными вкладками и фирменным стилем.
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # Стили данных
    data_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    highlight_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Amber 100

    # -------------------------------------------------------------------------
    # 1. Вкладка: «Анкета»
    # -------------------------------------------------------------------------
    ws_profile = wb.active
    ws_profile.title = "Анкета"
    ws_profile.views.sheetView[0].showGridLines = True

    # Заголовок страницы
    ws_profile.merge_cells("A1:D1")
    title_cell = ws_profile["A1"]
    title_cell.value = f"АНКЕТНАЯ КАРТОЧКА СУБЪЕКТА: {dossier.subject.full_name.upper()}"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(vertical="center")
    ws_profile.row_dimensions[1].height = 30

    profile_data = [
        ("ФИО субъекта", dossier.subject.full_name),
        ("ИИН / ИНН", dossier.subject.iin or "Не указан"),
        ("Дата рождения", dossier.subject.birth_date or "Не указана"),
        ("Место рождения", dossier.subject.birth_place or "Не указано"),
        ("Гражданство", dossier.subject.citizenship or "Не указано"),
        ("Пол", dossier.subject.gender or "Не указан"),
    ]

    _style_header_row(ws_profile, 3, 2, bg_hex="1E293B")
    ws_profile.cell(row=3, column=1, value="Параметр")
    ws_profile.cell(row=3, column=2, value="Значение")

    row_num = 4
    for label, val in profile_data:
        c1 = ws_profile.cell(row=row_num, column=1, value=label)
        c2 = ws_profile.cell(row=row_num, column=2, value=val)
        c1.font = bold_font
        c1.border = border_thin
        c1.fill = zebra_fill
        c2.font = data_font
        c2.border = border_thin
        row_num += 1

    # Резюме аналитика
    row_num += 1
    ws_profile.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=2)
    sum_header = ws_profile.cell(row=row_num, column=1, value="Сводные выводы аналитика")
    _style_header_row(ws_profile, row_num, 2, bg_hex="2563EB")
    row_num += 1

    ws_profile.merge_cells(start_row=row_num, start_column=1, end_row=row_num + 3, end_column=2)
    sum_body = ws_profile.cell(row=row_num, column=1, value=dossier.executive_summary)
    sum_body.font = Font(name="Arial", size=10, italic=True)
    sum_body.alignment = Alignment(vertical="top", wrap_text=True)
    sum_body.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    sum_body.border = border_thin

    _auto_fit_columns(ws_profile)

    # -------------------------------------------------------------------------
    # 2. Вкладка: «Контакты и адреса»
    # -------------------------------------------------------------------------
    ws_contacts = wb.create_sheet(title="Контакты и адреса")
    ws_contacts.views.sheetView[0].showGridLines = True

    # Блок контактов
    ws_contacts.cell(row=1, column=1, value="Средства связи субъекта").font = Font(name="Arial", size=12, bold=True, color="1E3A8A")
    _style_header_row(ws_contacts, 2, 2, bg_hex="1E293B")
    ws_contacts.cell(row=2, column=1, value="Тип контакта")
    ws_contacts.cell(row=2, column=2, value="Значение (список)")

    contact_rows = [
        ("Телефонные номера", ", ".join(dossier.contacts.phone_numbers) if dossier.contacts.phone_numbers else "Не выявлены"),
        ("Email адреса", ", ".join(dossier.contacts.emails) if dossier.contacts.emails else "Не выявлены"),
        ("Мессенджеры / Соцсети", ", ".join(dossier.contacts.social_profiles) if dossier.contacts.social_profiles else "Не выявлены"),
    ]
    cur_r = 3
    for label, val in contact_rows:
        c1 = ws_contacts.cell(row=cur_r, column=1, value=label)
        c2 = ws_contacts.cell(row=cur_r, column=2, value=val)
        c1.font = bold_font
        c1.fill = zebra_fill
        c1.border = border_thin
        c2.font = data_font
        c2.border = border_thin
        cur_r += 1

    # Блок адресов
    cur_r += 2
    ws_contacts.cell(row=cur_r, column=1, value="Реестр выявленных адресов").font = Font(name="Arial", size=12, bold=True, color="1E3A8A")
    cur_r += 1
    _style_header_row(ws_contacts, cur_r, 4, bg_hex="1E293B")
    ws_contacts.cell(row=cur_r, column=1, value="Тип адреса")
    ws_contacts.cell(row=cur_r, column=2, value="Город / Населенный пункт")
    ws_contacts.cell(row=cur_r, column=3, value="Регион / Область")
    ws_contacts.cell(row=cur_r, column=4, value="Полный адрес")

    cur_r += 1
    if dossier.addresses:
        for idx, addr in enumerate(dossier.addresses):
            fill = zebra_fill if idx % 2 == 1 else None
            for col_idx, val in enumerate([addr.address_type, addr.city or "—", addr.region or "—", addr.full_address], 1):
                cell = ws_contacts.cell(row=cur_r, column=col_idx, value=val)
                cell.font = data_font
                cell.border = border_thin
                if fill:
                    cell.fill = fill
            cur_r += 1
    else:
        ws_contacts.cell(row=cur_r, column=1, value="Адреса в документах не выявлены").font = Font(italic=True)

    _auto_fit_columns(ws_contacts)

    # -------------------------------------------------------------------------
    # 3. Вкладка: «Родственные связи»
    # -------------------------------------------------------------------------
    ws_rel = wb.create_sheet(title="Родственные связи")
    ws_rel.views.sheetView[0].showGridLines = True

    ws_rel.cell(row=1, column=1, value=f"Карта связей и аффилированных лиц для: {dossier.subject.full_name}").font = Font(name="Arial", size=12, bold=True, color="1E3A8A")
    _style_header_row(ws_rel, 2, 6, bg_hex="1E293B")
    ws_rel.cell(row=2, column=1, value="ФИО связанного лица")
    ws_rel.cell(row=2, column=2, value="Степень связи / родства")
    ws_rel.cell(row=2, column=3, value="ИИН / ИНН")
    ws_rel.cell(row=2, column=4, value="Дата рождения")
    ws_rel.cell(row=2, column=5, value="Общие признаки связи")
    ws_rel.cell(row=2, column=6, value="Примечания аналитика")

    cur_r = 3
    if dossier.relatives_and_affiliates:
        for idx, rel in enumerate(dossier.relatives_and_affiliates):
            shared_str = "; ".join(rel.shared_attributes) if rel.shared_attributes else ""
            fill = highlight_fill if rel.shared_attributes else (zebra_fill if idx % 2 == 1 else None)

            vals = [
                rel.full_name,
                rel.relation_type,
                rel.iin or "—",
                rel.birth_date or "—",
                shared_str or "—",
                rel.notes or "—"
            ]
            for col_idx, val in enumerate(vals, 1):
                cell = ws_rel.cell(row=cur_r, column=col_idx, value=val)
                cell.font = bold_font if (col_idx == 5 and rel.shared_attributes) else data_font
                cell.border = border_thin
                if fill:
                    cell.fill = fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            cur_r += 1
    else:
        ws_rel.cell(row=cur_r, column=1, value="Родственные связи в документах не выявлены").font = Font(italic=True)

    _auto_fit_columns(ws_rel)

    # -------------------------------------------------------------------------
    # 4. Вкладка: «Места работы»
    # -------------------------------------------------------------------------
    ws_emp = wb.create_sheet(title="Места работы")
    ws_emp.views.sheetView[0].showGridLines = True

    ws_emp.cell(row=1, column=1, value=f"Трудовая деятельность и аффилированный бизнес: {dossier.subject.full_name}").font = Font(name="Arial", size=12, bold=True, color="1E3A8A")
    _style_header_row(ws_emp, 2, 5, bg_hex="1E293B")
    ws_emp.cell(row=2, column=1, value="Организация / Работодатель")
    ws_emp.cell(row=2, column=2, value="БИН / ИНН компании")
    ws_emp.cell(row=2, column=3, value="Должность / Роль")
    ws_emp.cell(row=2, column=4, value="Период работы")
    ws_emp.cell(row=2, column=5, value="Дополнительные сведения")

    cur_r = 3
    if dossier.employment_history:
        for idx, emp in enumerate(dossier.employment_history):
            fill = zebra_fill if idx % 2 == 1 else None
            vals = [
                emp.organization,
                emp.bin or "—",
                emp.position or "—",
                emp.period or "—",
                emp.notes or "—"
            ]
            for col_idx, val in enumerate(vals, 1):
                cell = ws_emp.cell(row=cur_r, column=col_idx, value=val)
                cell.font = data_font
                cell.border = border_thin
                if fill:
                    cell.fill = fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            cur_r += 1
    else:
        ws_emp.cell(row=cur_r, column=1, value="Сведения о местах работы в документах не выявлены").font = Font(italic=True)

    _auto_fit_columns(ws_emp)

    wb.save(str(path))
    return path
